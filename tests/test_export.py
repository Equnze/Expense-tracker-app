import csv
from datetime import date, datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from PySide6.QtWidgets import QFileDialog, QMessageBox

from expense_tracker.services.exporter import (
    category_summary,
    export_csv,
    export_expenses,
    export_xlsx,
)


@pytest.fixture
def sample(expenses, categories, make_expense):
    make_expense(1299, date(2026, 10, 5), "Dining Out", "Lunch, with Ada", "Cash")
    make_expense(125000, date(2026, 10, 1), "Rent/Mortgage", "Rent", "Bank Transfer")
    make_expense(701, date(2026, 10, 2), "Dining Out", 'Café "Zoë"', "Cash", "treat,work")
    return expenses.list(), {c.id: c for c in categories.list()}


def test_category_summary_is_largest_first(sample):
    rows, by_id = sample
    assert category_summary(rows, by_id) == [("Rent/Mortgage", 125000), ("Dining Out", 2000)]
    assert category_summary([], by_id) == []


def test_csv_export(sample, tmp_path):
    rows, by_id = sample
    path = tmp_path / "out.csv"
    export_csv(path, rows, by_id)

    assert path.read_bytes().startswith(b"\xef\xbb\xbf")  # BOM so Excel reads UTF-8
    with open(path, newline="", encoding="utf-8-sig") as handle:
        lines = list(csv.reader(handle))
    assert lines[0] == ["Date", "Category", "Description", "Payment Method", "Tags", "Amount"]
    assert lines[1] == ["2026-10-05", "Dining Out", "Lunch, with Ada", "Cash", "", "12.99"]
    assert lines[2] == ["2026-10-02", "Dining Out", 'Café "Zoë"', "Cash", "treat,work", "7.01"]
    assert lines[3][5] == "1250.00"
    assert len(lines) == 4


def test_xlsx_export_has_expenses_and_summary_sheets(sample, tmp_path):
    rows, by_id = sample
    path = tmp_path / "out.xlsx"
    export_xlsx(path, rows, by_id, "£")

    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Expenses", "Summary"]

    sheet = workbook["Expenses"]
    values = list(sheet.iter_rows(values_only=True))
    assert values[0] == ("Date", "Category", "Description", "Payment Method", "Tags", "Amount")
    assert values[1] == (
        datetime(2026, 10, 5),
        "Dining Out",
        "Lunch, with Ada",
        "Cash",
        None,
        12.99,
    )
    assert len(values) == 4
    assert sheet["F2"].number_format == '"£"#,##0.00'
    assert sheet["A2"].number_format == "yyyy-mm-dd"
    assert sheet["A1"].font.bold
    # Amounts are real numbers that add up exactly in Excel.
    assert sum(Decimal(str(row[5])) for row in values[1:]) == Decimal("1270.00")

    summary = list(workbook["Summary"].iter_rows(values_only=True))
    assert summary[0] == ("Category", "Amount", "% of total")
    assert summary[1][:2] == ("Rent/Mortgage", 1250)
    assert summary[1][2] == pytest.approx(125000 / 127000)
    assert summary[2][:2] == ("Dining Out", 20)
    assert summary[3] == ("Total", 1270, 1)
    assert workbook["Summary"]["C2"].number_format == "0.0%"


def test_xlsx_export_of_nothing_is_still_a_valid_workbook(tmp_path, categories):
    path = tmp_path / "empty.xlsx"
    export_xlsx(path, [], {c.id: c for c in categories.list()})
    workbook = load_workbook(path)
    assert workbook["Expenses"].max_row == 1
    assert list(workbook["Summary"].iter_rows(values_only=True))[1] == ("Total", 0, 0)


def test_format_follows_the_file_extension(sample, tmp_path):
    rows, by_id = sample
    export_expenses(tmp_path / "a.XLSX", rows, by_id)
    export_expenses(tmp_path / "b.csv", rows, by_id)
    assert load_workbook(tmp_path / "a.XLSX").sheetnames == ["Expenses", "Summary"]
    assert (tmp_path / "b.csv").read_text(encoding="utf-8-sig").startswith("Date,")


# --- UI ---------------------------------------------------------------------


def test_export_button_writes_only_the_filtered_rows(
    window, ctx, make_expense, statuses, tmp_path, monkeypatch
):
    tab = window.expenses_tab
    assert not tab.export_button.isEnabled()

    today = date.today()
    make_expense(100, today, "Travel", "Bus")
    make_expense(200, today, "Gifts", "Flowers")
    make_expense(400, date(2001, 5, 5), "Gifts", "Old gift")
    ctx.data_changed.emit()
    tab.filters.search.setText("flow")
    assert tab.export_button.isEnabled()

    # The user types a name without an extension and picks the CSV file type.
    monkeypatch.setattr(
        QFileDialog,
        "getSaveFileName",
        lambda *args, **kwargs: (str(tmp_path / "mine"), "CSV file (*.csv)"),
    )
    tab.export_button.click()

    with open(tmp_path / "mine.csv", newline="", encoding="utf-8-sig") as handle:
        lines = list(csv.reader(handle))
    assert [line[2] for line in lines[1:]] == ["Flowers"]
    assert statuses[-1] == "Exported 1 row to mine.csv"
    assert ctx.settings.get("last_export_dir") == str(tmp_path)

    tab.filters.search.setText("")
    assert tab.export_to(tmp_path / "all.xlsx")
    assert load_workbook(tmp_path / "all.xlsx")["Expenses"].max_row == 3
    assert statuses[-1] == "Exported 2 rows to all.xlsx"


def test_export_cancelled_or_unwritable_is_handled(
    window, ctx, make_expense, statuses, tmp_path, monkeypatch
):
    make_expense(100, date.today())
    ctx.data_changed.emit()
    tab = window.expenses_tab

    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args, **kwargs: ("", ""))
    tab.export_button.click()
    assert not any(p.suffix in (".csv", ".xlsx") for p in tmp_path.iterdir())

    warnings = []
    monkeypatch.setattr(QMessageBox, "warning", lambda *args: warnings.append(args[2]))
    assert not tab.export_to(tmp_path / "missing-folder" / "x.csv")
    assert len(warnings) == 1 and "x.csv" in warnings[0]
    assert not any(message.startswith("Exported") for message in statuses)
