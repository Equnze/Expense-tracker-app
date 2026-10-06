"""Export expenses to CSV or Excel."""

from __future__ import annotations

import csv
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from expense_tracker.models import Category, Expense
from expense_tracker.utils.money import format_plain

HEADERS = ["Date", "Category", "Description", "Payment Method", "Tags", "Amount"]
SUMMARY_HEADERS = ["Category", "Amount", "% of total"]

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", start_color="EEF1F5")


def _category_name(expense: Expense, categories: dict[int, Category]) -> str:
    category = categories.get(expense.category_id)
    return category.name if category else ""


def category_summary(
    expenses: list[Expense], categories: dict[int, Category]
) -> list[tuple[str, int]]:
    """Return (category name, total cents) for the given expenses, largest first."""
    totals: dict[str, int] = {}
    for expense in expenses:
        name = _category_name(expense, categories)
        totals[name] = totals.get(name, 0) + expense.amount_cents
    return sorted(totals.items(), key=lambda item: (-item[1], item[0].casefold()))


def export_csv(
    path: Path | str, expenses: list[Expense], categories: dict[int, Category]
) -> None:
    """Write the expenses as CSV. Dates are ISO and amounts are plain decimals,
    so the file opens cleanly in Excel and can be read by other programs."""
    # utf-8-sig makes Excel recognise accented characters and currency symbols.
    with open(path, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADERS)
        for expense in expenses:
            writer.writerow(
                [
                    expense.date.isoformat(),
                    _category_name(expense, categories),
                    expense.description,
                    expense.payment_method,
                    expense.tags,
                    format_plain(expense.amount_cents),
                ]
            )


def _money_format(currency: str) -> str:
    symbol = currency.replace('"', "")
    return f'"{symbol}"#,##0.00' if symbol else "#,##0.00"


def _style_header(sheet, widths: list[int]) -> None:
    for column, width in enumerate(widths, start=1):
        cell = sheet.cell(row=1, column=column)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"


def export_xlsx(
    path: Path | str,
    expenses: list[Expense],
    categories: dict[int, Category],
    currency: str = "$",
) -> None:
    """Write an Excel workbook: an Expenses sheet and a per-category Summary sheet."""
    money_format = _money_format(currency)
    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Expenses"
    sheet.append(HEADERS)
    for expense in expenses:
        sheet.append(
            [
                expense.date,
                _category_name(expense, categories),
                expense.description,
                expense.payment_method,
                expense.tags,
                Decimal(expense.amount_cents) / 100,
            ]
        )
    for row in range(2, len(expenses) + 2):
        sheet.cell(row=row, column=1).number_format = "yyyy-mm-dd"
        sheet.cell(row=row, column=6).number_format = money_format
    sheet.cell(row=1, column=6).alignment = Alignment(horizontal="right")
    _style_header(sheet, [12, 22, 42, 18, 22, 14])
    if expenses:
        sheet.auto_filter.ref = f"A1:F{len(expenses) + 1}"

    summary = workbook.create_sheet("Summary")
    summary.append(SUMMARY_HEADERS)
    total = sum(expense.amount_cents for expense in expenses)
    for name, cents in category_summary(expenses, categories):
        summary.append([name, Decimal(cents) / 100, cents / total])
    summary.append(["Total", Decimal(total) / 100, 1 if total else 0])
    last_row = summary.max_row
    for row in range(2, last_row + 1):
        summary.cell(row=row, column=2).number_format = money_format
        summary.cell(row=row, column=3).number_format = "0.0%"
    for column in (1, 2, 3):
        summary.cell(row=last_row, column=column).font = _HEADER_FONT
    for column in (2, 3):
        summary.cell(row=1, column=column).alignment = Alignment(horizontal="right")
    _style_header(summary, [26, 16, 12])

    workbook.save(path)


def export_expenses(
    path: Path | str,
    expenses: list[Expense],
    categories: dict[int, Category],
    currency: str = "$",
) -> None:
    """Export to the format implied by the file extension (.xlsx or .csv)."""
    if Path(path).suffix.lower() == ".xlsx":
        export_xlsx(path, expenses, categories, currency)
    else:
        export_csv(path, expenses, categories)
